#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from PIL import Image, ImageDraw, ImageFont
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage

# Create a placeholder logo image
width, height = 400, 120
img = Image.new('RGB', (width, height), color='white')
draw = ImageDraw.Draw(img)

# Draw a border
border_color = (31, 78, 120)  # Dark blue
draw.rectangle([0, 0, width-1, height-1], outline=border_color, width=3)

# Try to use a nice font, fallback to default
try:
    font_large = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 28)
    font_small = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 16)
except:
    font_large = ImageFont.load_default()
    font_small = ImageFont.load_default()

# Add text
text1 = "CTP"
text2 = "CONSULTING ENGINEERS"

# Position text
draw.text((width//2, height//2 - 20), text1, fill=border_color, font=font_large, anchor="mm")
draw.text((width//2, height//2 + 20), text2, fill=(68, 114, 196), font=font_small, anchor="mm")

# Save the placeholder
logo_filename = "CTP_Logo_Placeholder.png"
img.save(logo_filename)
print(f"✓ Created placeholder logo: {logo_filename}")

# Insert into Excel template
try:
    wb = load_workbook("CTP_Document_Submittal_Response_DSR_Template.xlsx")
    ws = wb.active

    # Remove the placeholder text from A1
    ws['A1'] = ""

    # Add image to the header
    excel_img = XLImage(logo_filename)
    excel_img.width = 300
    excel_img.height = 90
    ws.add_image(excel_img, 'B1')

    wb.save("CTP_Document_Submittal_Response_DSR_Template.xlsx")
    print("✓ Placeholder logo inserted into Excel template!")
    print("\nNOTE: Replace this placeholder with the actual CTP logo from:")
    print("  - https://ctp-llp.com/ (UK)")
    print("  - https://www.c-t-p.co.za/ (South Africa)")

except Exception as e:
    print(f"✗ Could not insert logo into Excel: {e}")
