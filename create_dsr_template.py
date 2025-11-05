#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime

# Create workbook
wb = Workbook()
ws = wb.active
ws.title = "DSR Form"

# Define styles
header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
header_font = Font(bold=True, color="FFFFFF", size=12)
subheader_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
subheader_font = Font(bold=True, color="FFFFFF", size=11)
section_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
section_font = Font(bold=True, size=10)
label_font = Font(bold=True, size=10)
normal_font = Font(size=10)
center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
thin_border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

# Set column widths
ws.column_dimensions['A'].width = 3
ws.column_dimensions['B'].width = 25
ws.column_dimensions['C'].width = 35
ws.column_dimensions['D'].width = 25
ws.column_dimensions['E'].width = 35

# Row 1-3: Header with logo placeholder
ws.merge_cells('A1:E3')
ws['A1'] = "[INSERTAR LOGO CTP CONSULTING ENGINEERS AQUÍ]"
ws['A1'].font = Font(bold=True, size=14, italic=True)
ws['A1'].alignment = center_align
ws['A1'].fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
ws['A1'].border = thin_border
ws.row_dimensions[1].height = 25
ws.row_dimensions[2].height = 25
ws.row_dimensions[3].height = 25

# Row 4: Title
row = 4
ws.merge_cells(f'A{row}:E{row}')
ws[f'A{row}'] = "DOCUMENT SUBMITTAL RESPONSE / DESIGN SUBMITTAL REVIEW (DSR)"
ws[f'A{row}'].font = header_font
ws[f'A{row}'].fill = header_fill
ws[f'A{row}'].alignment = center_align
ws[f'A{row}'].border = thin_border
ws.row_dimensions[row].height = 30

# Row 5: Empty
row += 1
ws.row_dimensions[row].height = 5

# Row 6: PROJECT INFORMATION
row += 1
ws.merge_cells(f'B{row}:E{row}')
ws[f'B{row}'] = "PROJECT INFORMATION"
ws[f'B{row}'].font = subheader_font
ws[f'B{row}'].fill = subheader_fill
ws[f'B{row}'].alignment = center_align
ws[f'B{row}'].border = thin_border

# Project details
row += 1
ws[f'B{row}'] = "Project Name:"
ws[f'B{row}'].font = label_font
ws[f'B{row}'].border = thin_border
ws.merge_cells(f'C{row}:E{row}')
ws[f'C{row}'].border = thin_border
ws[f'C{row}'].alignment = left_align

row += 1
ws[f'B{row}'] = "Project Number:"
ws[f'B{row}'].font = label_font
ws[f'B{row}'].border = thin_border
ws[f'C{row}'].border = thin_border
ws[f'D{row}'] = "Date:"
ws[f'D{row}'].font = label_font
ws[f'D{row}'].border = thin_border
ws[f'E{row}'].border = thin_border

row += 1
ws[f'B{row}'] = "Client:"
ws[f'B{row}'].font = label_font
ws[f'B{row}'].border = thin_border
ws.merge_cells(f'C{row}:E{row}')
ws[f'C{row}'].border = thin_border

row += 1
ws[f'B{row}'] = "Contractor/Designer:"
ws[f'B{row}'].font = label_font
ws[f'B{row}'].border = thin_border
ws.merge_cells(f'C{row}:E{row}')
ws[f'C{row}'].border = thin_border

# Row: Empty
row += 1
ws.row_dimensions[row].height = 5

# SUBMITTAL INFORMATION
row += 1
ws.merge_cells(f'B{row}:E{row}')
ws[f'B{row}'] = "SUBMITTAL INFORMATION"
ws[f'B{row}'].font = subheader_font
ws[f'B{row}'].fill = subheader_fill
ws[f'B{row}'].alignment = center_align
ws[f'B{row}'].border = thin_border

row += 1
ws[f'B{row}'] = "Submittal Number:"
ws[f'B{row}'].font = label_font
ws[f'B{row}'].border = thin_border
ws[f'C{row}'].border = thin_border
ws[f'D{row}'] = "Revision:"
ws[f'D{row}'].font = label_font
ws[f'D{row}'].border = thin_border
ws[f'E{row}'].border = thin_border

row += 1
ws[f'B{row}'] = "Document Title:"
ws[f'B{row}'].font = label_font
ws[f'B{row}'].border = thin_border
ws.merge_cells(f'C{row}:E{row}')
ws[f'C{row}'].border = thin_border

row += 1
ws[f'B{row}'] = "Document Number/Drawing No.:"
ws[f'B{row}'].font = label_font
ws[f'B{row}'].border = thin_border
ws.merge_cells(f'C{row}:E{row}')
ws[f'C{row}'].border = thin_border

row += 1
ws[f'B{row}'] = "Date Received:"
ws[f'B{row}'].font = label_font
ws[f'B{row}'].border = thin_border
ws[f'C{row}'].border = thin_border
ws[f'D{row}'] = "Date Reviewed:"
ws[f'D{row}'].font = label_font
ws[f'D{row}'].border = thin_border
ws[f'E{row}'].border = thin_border

# Row: Empty
row += 1
ws.row_dimensions[row].height = 5

# REVIEW STATUS
row += 1
ws.merge_cells(f'B{row}:E{row}')
ws[f'B{row}'] = "REVIEW STATUS"
ws[f'B{row}'].font = subheader_font
ws[f'B{row}'].fill = subheader_fill
ws[f'B{row}'].alignment = center_align
ws[f'B{row}'].border = thin_border

row += 1
ws[f'B{row}'] = "☐ APPROVED"
ws[f'B{row}'].font = Font(bold=True, size=11)
ws[f'B{row}'].border = thin_border
ws.merge_cells(f'C{row}:E{row}')
ws[f'C{row}'] = "No exceptions taken. Proceed with work."
ws[f'C{row}'].alignment = left_align
ws[f'C{row}'].border = thin_border
ws.row_dimensions[row].height = 20

row += 1
ws[f'B{row}'] = "☐ APPROVED AS NOTED"
ws[f'B{row}'].font = Font(bold=True, size=11)
ws[f'B{row}'].border = thin_border
ws.merge_cells(f'C{row}:E{row}')
ws[f'C{row}'] = "Proceed with work. Incorporate noted corrections."
ws[f'C{row}'].alignment = left_align
ws[f'C{row}'].border = thin_border
ws.row_dimensions[row].height = 20

row += 1
ws[f'B{row}'] = "☐ REVISE AND RESUBMIT"
ws[f'B{row}'].font = Font(bold=True, size=11)
ws[f'B{row}'].border = thin_border
ws.merge_cells(f'C{row}:E{row}')
ws[f'C{row}'] = "Do not proceed. Make corrections and resubmit."
ws[f'C{row}'].alignment = left_align
ws[f'C{row}'].border = thin_border
ws.row_dimensions[row].height = 20

row += 1
ws[f'B{row}'] = "☐ REJECTED"
ws[f'B{row}'].font = Font(bold=True, size=11)
ws[f'B{row}'].border = thin_border
ws.merge_cells(f'C{row}:E{row}')
ws[f'C{row}'] = "Does not meet requirements. Complete redesign required."
ws[f'C{row}'].alignment = left_align
ws[f'C{row}'].border = thin_border
ws.row_dimensions[row].height = 20

# Row: Empty
row += 1
ws.row_dimensions[row].height = 5

# STRUCTURAL DESIGN REVIEW CHECKLIST
row += 1
ws.merge_cells(f'B{row}:E{row}')
ws[f'B{row}'] = "STRUCTURAL DESIGN REVIEW CHECKLIST"
ws[f'B{row}'].font = subheader_font
ws[f'B{row}'].fill = subheader_fill
ws[f'B{row}'].alignment = center_align
ws[f'B{row}'].border = thin_border

# Checklist headers
row += 1
ws[f'B{row}'] = "Item"
ws[f'B{row}'].font = label_font
ws[f'B{row}'].fill = section_fill
ws[f'B{row}'].alignment = center_align
ws[f'B{row}'].border = thin_border

ws.merge_cells(f'C{row}:D{row}')
ws[f'C{row}'] = "Description"
ws[f'C{row}'].font = label_font
ws[f'C{row}'].fill = section_fill
ws[f'C{row}'].alignment = center_align
ws[f'C{row}'].border = thin_border

ws[f'E{row}'] = "Status"
ws[f'E{row}'].font = label_font
ws[f'E{row}'].fill = section_fill
ws[f'E{row}'].alignment = center_align
ws[f'E{row}'].border = thin_border

# Checklist items
checklist_items = [
    ("1", "Design calculations provided and complete"),
    ("2", "Load combinations comply with applicable codes"),
    ("3", "Material specifications clearly identified"),
    ("4", "Structural drawings are clear and coordinated"),
    ("5", "Design meets code requirements (IBC, ACI, AISC, etc.)"),
    ("6", "Lateral load resisting system adequately designed"),
    ("7", "Foundation design adequate for soil conditions"),
    ("8", "Connection details are adequate and constructible"),
    ("9", "Deflection limits are met"),
    ("10", "Special inspections requirements identified"),
    ("11", "Seismic design criteria properly applied"),
    ("12", "Wind load calculations verified"),
    ("13", "Construction details are practical and buildable"),
    ("14", "All structural elements properly sized"),
    ("15", "Reinforcement details meet code requirements"),
]

for item_num, item_desc in checklist_items:
    row += 1
    ws[f'B{row}'] = item_num
    ws[f'B{row}'].alignment = center_align
    ws[f'B{row}'].border = thin_border

    ws.merge_cells(f'C{row}:D{row}')
    ws[f'C{row}'] = item_desc
    ws[f'C{row}'].alignment = left_align
    ws[f'C{row}'].border = thin_border

    ws[f'E{row}'].border = thin_border
    ws.row_dimensions[row].height = 25

# Row: Empty
row += 1
ws.row_dimensions[row].height = 5

# COMMENTS AND OBSERVATIONS
row += 1
ws.merge_cells(f'B{row}:E{row}')
ws[f'B{row}'] = "COMMENTS AND OBSERVATIONS"
ws[f'B{row}'].font = subheader_font
ws[f'B{row}'].fill = subheader_fill
ws[f'B{row}'].alignment = center_align
ws[f'B{row}'].border = thin_border

row += 1
ws.merge_cells(f'B{row}:E{row + 8}')
ws[f'B{row}'].border = thin_border
ws[f'B{row}'].alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
for i in range(9):
    ws.row_dimensions[row + i].height = 25

row += 9

# Row: Empty
row += 1
ws.row_dimensions[row].height = 5

# REVIEWER INFORMATION
row += 1
ws.merge_cells(f'B{row}:E{row}')
ws[f'B{row}'] = "REVIEWER INFORMATION"
ws[f'B{row}'].font = subheader_font
ws[f'B{row}'].fill = subheader_fill
ws[f'B{row}'].alignment = center_align
ws[f'B{row}'].border = thin_border

row += 1
ws[f'B{row}'] = "Reviewed By:"
ws[f'B{row}'].font = label_font
ws[f'B{row}'].border = thin_border
ws.merge_cells(f'C{row}:E{row}')
ws[f'C{row}'].border = thin_border
ws.row_dimensions[row].height = 20

row += 1
ws[f'B{row}'] = "Title:"
ws[f'B{row}'].font = label_font
ws[f'B{row}'].border = thin_border
ws[f'C{row}'].border = thin_border
ws[f'D{row}'] = "License No.:"
ws[f'D{row}'].font = label_font
ws[f'D{row}'].border = thin_border
ws[f'E{row}'].border = thin_border
ws.row_dimensions[row].height = 20

row += 1
ws[f'B{row}'] = "Signature:"
ws[f'B{row}'].font = label_font
ws[f'B{row}'].border = thin_border
ws.merge_cells(f'C{row}:D{row}')
ws[f'C{row}'].border = thin_border
ws[f'E{row}'] = "Date:"
ws[f'E{row}'].font = label_font
ws[f'E{row}'].border = thin_border
ws.row_dimensions[row].height = 30

# Row: Empty
row += 1
ws.row_dimensions[row].height = 5

# FOOTER
row += 1
ws.merge_cells(f'B{row}:E{row}')
ws[f'B{row}'] = "CTP Consulting Engineers"
ws[f'B{row}'].font = Font(bold=True, size=9, color="666666")
ws[f'B{row}'].alignment = center_align
ws.row_dimensions[row].height = 15

row += 1
ws.merge_cells(f'B{row}:E{row}')
ws[f'B{row}'] = "Structural Engineering Services"
ws[f'B{row}'].font = Font(size=9, color="666666")
ws[f'B{row}'].alignment = center_align
ws.row_dimensions[row].height = 15

# Save the workbook
filename = "CTP_Document_Submittal_Response_DSR_Template.xlsx"
wb.save(filename)
print(f"Template created successfully: {filename}")
print(f"Total rows: {row}")
